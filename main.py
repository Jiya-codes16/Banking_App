import streamlit as st
import openpyxl
import os
import uuid
from datetime import datetime


# ==========================================================================================
# INITIALIZATION
# ==========================================================================================

file_name = "bank_records.xlsx"


# ==========================================================================================
# PAGE CONFIGURATION
# ==========================================================================================

st.set_page_config(
    page_title="Trust Bank",
    page_icon="🏦",
    layout="wide"
)


# ==========================================================================================
# CUSTOM BANK UI
# ==========================================================================================

st.markdown("""
<style>

.stApp {
    background:
        radial-gradient(circle at 10% 20%, rgba(99, 102, 241, 0.35), transparent 30%),
        radial-gradient(circle at 90% 10%, rgba(236, 72, 153, 0.30), transparent 30%),
        radial-gradient(circle at 50% 90%, rgba(14, 165, 233, 0.25), transparent 35%),
        linear-gradient(135deg, #080b18, #11152b, #080b18);
    color: white;
}

/* Moving Trust Bank Title */

.bank-title {
    text-align: center;
    font-size: 65px;
    font-weight: 900;
    letter-spacing: 8px;
    margin-top: 10px;
    margin-bottom: 5px;

    background: linear-gradient(
        90deg,
        #60a5fa,
        #c084fc,
        #f472b6,
        #38bdf8,
        #60a5fa
    );

    background-size: 300% auto;

    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;

    animation: bankGradient 5s linear infinite;
}

/* Moving gradient */

@keyframes bankGradient {

    0% {
        background-position: 0% center;
    }

    50% {
        background-position: 100% center;
    }

    100% {
        background-position: 0% center;
    }
}

/* Subtitle */

.bank-subtitle {
    text-align: center;
    font-size: 21px;
    color: #dbeafe;
    margin-bottom: 25px;
    animation: fadeIn 2s ease-in-out;
}

@keyframes fadeIn {

    from {
        opacity: 0;
        transform: translateY(10px);
    }

    to {
        opacity: 1;
        transform: translateY(0);
    }
}


/* Glass cards */

div[data-testid="stMetric"] {

    background: rgba(255, 255, 255, 0.08);
    border: 1px solid rgba(255, 255, 255, 0.15);
    border-radius: 18px;
    padding: 18px;

    backdrop-filter: blur(12px);

    box-shadow:
        0 8px 30px rgba(0, 0, 0, 0.25);
}


/* Buttons */

.stButton > button {

    width: 100%;

    border-radius: 12px;

    border: 1px solid rgba(255,255,255,0.2);

    background: linear-gradient(
        135deg,
        rgba(99,102,241,0.8),
        rgba(236,72,153,0.8)
    );

    color: white;

    font-weight: 700;

    padding: 10px;

    transition: all 0.3s ease;
}


.stButton > button:hover {

    transform: translateY(-3px) scale(1.02);

    box-shadow:
        0 8px 25px rgba(236,72,153,0.35);
}


/* Input boxes */

div[data-baseweb="input"] {

    background: rgba(255,255,255,0.08);

    border-radius: 10px;

}


/* Sidebar */

section[data-testid="stSidebar"] {

    background:
        linear-gradient(
            180deg,
            #0b1026,
            #16122e,
            #0b1026
        );

    border-right: 1px solid rgba(255,255,255,0.1);
}


/* Headers */

h1, h2, h3 {

    color: #f8fafc !important;

}


/* Divider */

hr {

    border-color: rgba(255,255,255,0.15);

}

</style>
""", unsafe_allow_html=True)


# ==========================================================================================
# ANIMATED BANK TITLE
# ==========================================================================================

st.markdown(
    '<div class="bank-title">🏦 TRUST BANK</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="bank-subtitle">'
    '💳 Secure Banking  💰 Smart Money Management  📊 Easy Transactions'
    '</div>',
    unsafe_allow_html=True
)

st.divider()

# ==========================================================================================
# CURRENT DATE AND TIME
# ==========================================================================================

def cur_date():
    now = datetime.now()
    formatted_1 = now.strftime("%Y-%m-%d %H-%M-%S")
    return formatted_1


# ==========================================================================================
# EXCEL CREATION FUNCTION
# ==========================================================================================

def create_excel_file():
    """This Function is used to create an excel file."""

    if not os.path.exists(file_name):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Bank Records"

        headers = [
            "Accout No.",
            "Name",
            "PIN",
            "Transaction ID",
            "Transaction Type",
            "Amount",
            "Previous Balance",
            "Current Balance",
            "Date-Time"
        ]

        sheet.append(headers)
        wb.save(file_name)
        wb.close()


# ==========================================================================================
# ERROR MESSAGE
# ==========================================================================================

def err_msg(error):
    st.error(f"❌ Invalid {error} !!! Please Try Again...")


# ==========================================================================================
# CHECK WHETHER ACCOUNT ALREADY EXISTS
# ==========================================================================================

def account_exists(acc_nr):

    wb = openpyxl.load_workbook(file_name)
    sheet = wb.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(acc_nr):
            found = True
            break

    wb.close()

    return found


# ==========================================================================================
# CREATE ACCOUNT
# ==========================================================================================

def create_account():

    st.header("👤 Create Account")

    name = st.text_input(
        "✏️ Enter Your Name",
        key="create_name"
    )

    acc_nr = st.text_input(
        "💳 Enter Your Account Number",
        key="create_account"
    )

    pin = st.text_input(
        "🔑 Create a 4-digit PIN",
        type="password",
        max_chars=4,
        key="create_pin"
    )

    amount = st.number_input(
        "💰 Enter Initial Amount",
        min_value=0.0,
        step=100.0,
        key="create_amount"
    )

    if st.button("Create Account", type="primary"):

        if name.strip() == "":
            st.error("❌ Please enter your name.")
            return

        if not acc_nr.isdigit():
            st.error("❌ Invalid Account Number.")
            return

        if account_exists(acc_nr):
            st.error("⚠️ This Account Number is Already Taken!")
            return

        if len(pin) != 4 or not pin.isdigit():
            st.error("❌ PIN must contain exactly 4 digits.")
            return

        if amount < 0:
            err_msg("Amount")
            return

        transaction_id = str(uuid.uuid4())[:8]

        status = "Opening"

        dt = cur_date()

        wb = openpyxl.load_workbook(file_name)
        sheet = wb.active

        sheet.append([
            acc_nr,
            name,
            pin,
            transaction_id,
            status,
            amount,
            0,
            amount,
            dt
        ])

        wb.save(file_name)
        wb.close()

        st.success("✅ Account Created Successfully!")

        st.info(
            f"Account Number: {acc_nr}\n\n"
            f"Initial Balance: ₹{amount:.2f}"
        )


# ==========================================================================================
# GET CURRENT BALANCE
# ==========================================================================================

def get_balance(acc_nr):

    wb = openpyxl.load_workbook(file_name)
    sheet = wb.active

    current_balance = None
    account_found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(acc_nr):

            current_balance = row[7]
            account_found = True

    wb.close()

    if account_found:
        return current_balance

    else:
        return None


# ==========================================================================================
# GET USER NAME AND PIN
# ==========================================================================================

def get_user_details(acc_nr):

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    name = None
    pin = None

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(acc_nr):

            name = row[1]
            pin = row[2]

            break

    workbook.close()

    return name, pin


# ==========================================================================================
# CHECK BALANCE
# ==========================================================================================

def check_balance(acc_nr):

    balance = get_balance(acc_nr)

    if balance is not None:

        st.metric(
            "💰 Current Balance",
            f"₹{balance:.2f}"
        )

    else:
        st.error("❌ Account Details Error!")


# ==========================================================================================
# DEPOSIT MONEY
# ==========================================================================================

def deposit(acc_nr):

    st.subheader("💰 Deposit Money")

    amount = st.number_input(
        "Enter Deposit Amount",
        min_value=0.0,
        step=100.0,
        key="deposit_amount"
    )

    if st.button("Deposit Money", type="primary"):

        if amount <= 0:
            st.error("❌ Amount Must Be Greater Than 0!")
            return

        previous_balance = get_balance(acc_nr)

        if previous_balance is None:
            st.error("❌ Account not found.")
            return

        new_balance = previous_balance + amount

        name, pin = get_user_details(acc_nr)

        if name is None:
            st.error(
                "❌ Seems Data Is Corrupted. "
                "Name And PIN Doesn't Exist But Account Does!"
            )
            return

        transaction_id = str(uuid.uuid4())[:8]

        date_time = datetime.now().strftime(
            "%d-%m-%Y %H:%M:%S"
        )

        wb = openpyxl.load_workbook(file_name)
        sheet = wb.active

        sheet.append([
            acc_nr,
            name,
            pin,
            transaction_id,
            "Deposit",
            amount,
            previous_balance,
            new_balance,
            date_time
        ])

        wb.save(file_name)
        wb.close()

        st.success(
            f"✅ ₹{amount:.2f} Deposited Successfully!"
        )

        st.metric(
            "New Balance",
            f"₹{new_balance:.2f}"
        )


# ==========================================================================================
# WITHDRAW MONEY
# ==========================================================================================

def withdraw(acc_nr):

    st.subheader("💳 Withdraw Money")

    amount = st.number_input(
        "Enter Withdraw Amount",
        min_value=0.0,
        step=100.0,
        key="withdraw_amount"
    )

    if st.button("Withdraw Money", type="primary"):

        if amount <= 0:
            st.error("❌ Amount Must Be Greater Than 0!")
            return

        previous_balance = get_balance(acc_nr)

        if previous_balance is None:
            st.error("❌ Account not found.")
            return

        if amount > previous_balance:
            st.error("⚠️ Insufficient Balance!")
            return

        new_balance = previous_balance - amount

        name, pin = get_user_details(acc_nr)

        if name is None:
            st.error(
                "❌ Seems Data Is Corrupted. "
                "Name And PIN Doesn't Exist But Account Does!"
            )
            return

        transaction_id = str(uuid.uuid4())[:8]

        date_time = datetime.now().strftime(
            "%d-%m-%Y %H:%M:%S"
        )

        wb = openpyxl.load_workbook(file_name)
        sheet = wb.active

        sheet.append([
            acc_nr,
            name,
            pin,
            transaction_id,
            "Withdraw",
            amount,
            previous_balance,
            new_balance,
            date_time
        ])

        wb.save(file_name)
        wb.close()

        st.success(
            f"✅ ₹{amount:.2f} Withdrawn Successfully!"
        )

        st.metric(
            "New Balance",
            f"₹{new_balance:.2f}"
        )


# ==========================================================================================
# TRANSACTION HISTORY
# ==========================================================================================

def transaction_history(acc_nr):

    st.subheader("📋 Transaction History")

    wb = openpyxl.load_workbook(file_name)
    sheet = wb.active

    transactions = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == str(acc_nr):

            transactions.append({
                "Transaction ID": str(row[3]),
                "Type": str(row[4]),
                "Amount": f"₹{row[5]:.2f}",
                "Previous Balance": f"₹{row[6]:.2f}",
                "Current Balance": f"₹{row[7]:.2f}",
                "Date": str(row[8])
            })

    wb.close()

    if transactions:

        st.dataframe(
            transactions,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.warning("⚠️ No Transaction Found!")


# ==========================================================================================
# LOGIN
# ==========================================================================================

def login():

    st.header("🔐 Login")

    acc_nr = st.text_input(
        "💳 Enter Your Account Number",
        key="login_account"
    )

    pin = st.text_input(
        "🔑 Enter Your 4-digit PIN",
        type="password",
        max_chars=4,
        key="login_pin"
    )

    if st.button("Login", type="primary"):

        if not acc_nr or not pin:
            st.error("❌ Please enter Account Number and PIN.")
            return

        wb = openpyxl.load_workbook(file_name)
        sheet = wb.active

        login_success = False
        user_name = ""

        for row in sheet.iter_rows(
            min_row=2,
            max_col=3,
            values_only=True
        ):

            if (
                str(row[0]) == str(acc_nr)
                and str(row[2]) == str(pin)
            ):

                user_name = row[1]
                login_success = True
                break

        wb.close()

        if not login_success:

            st.error("❌ Invalid Account Number Or PIN.")
            return

        st.session_state.logged_in = True
        st.session_state.account_number = acc_nr
        st.session_state.user_name = user_name

        st.success("✅ LOGIN SUCCESSFUL!")

        st.rerun()


# ==========================================================================================
# BANKING DASHBOARD
# ==========================================================================================

def sub_menu(acc_nr):

    user_name = st.session_state.user_name

    st.success(f"Welcome, {user_name}! 👋")

    balance = get_balance(acc_nr)

    # Dashboard boxes
    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric(
            "👤 Account Holder",
            user_name
        )

    with col2:
        st.metric(
            "💰 Current Balance",
            f"₹{balance:.2f}"
        )

    with col3:
        st.metric(
            "💳 Account Number",
            acc_nr
        )

    st.divider()

    option = st.sidebar.radio(
        "🏦 Banking Options",
        [
            "Dashboard",
            "Check Balance",
            "Deposit Money",
            "Withdraw Money",
            "Transaction History",
            "Logout"
        ]
    )

    if option == "Dashboard":

        st.header("🏦 Banking Dashboard")

        st.write(
            "Manage your Trust Bank account easily."
        )

        col1, col2 = st.columns(2)

        with col1:

            st.info(
                "💰 Use *Check Balance* "
                "to view your current account balance."
            )

            st.info(
                "💵 Use *Deposit Money* "
                "to add money to your account."
            )

        with col2:

            st.info(
                "💳 Use *Withdraw Money* "
                "to withdraw money."
            )

            st.info(
                "📋 Use *Transaction History* "
                "to view your transactions."
            )

    elif option == "Check Balance":

        st.header("💰 Check Balance")

        check_balance(acc_nr)

    elif option == "Deposit Money":

        deposit(acc_nr)

    elif option == "Withdraw Money":

        withdraw(acc_nr)

    elif option == "Transaction History":

        transaction_history(acc_nr)

    elif option == "Logout":

        st.session_state.logged_in = False
        st.session_state.account_number = None
        st.session_state.user_name = None

        st.success("👋 Logged out successfully!")

        st.rerun()


# ==========================================================================================
# INITIALIZE SESSION STATE
# ==========================================================================================

if "logged_in" not in st.session_state:

    st.session_state.logged_in = False

if "account_number" not in st.session_state:

    st.session_state.account_number = None

if "user_name" not in st.session_state:

    st.session_state.user_name = None


# ==========================================================================================
# MAIN PROGRAM
# ==========================================================================================

create_excel_file()


# ==========================================================================================
# SIDEBAR
# ==========================================================================================

if not st.session_state.logged_in:

    st.sidebar.title("🏦 TRUST BANK")

    main_option = st.sidebar.radio(
        "Select Option",
        [
            "Home",
            "Create Account",
            "Login"
        ]
    )

    if main_option == "Home":

        st.header("🏦 Welcome to Trust Bank")

        st.write(
            "A simple and secure banking management system."
        )

        st.divider()

        col1, col2, col3 = st.columns(3)

        with col1:
            st.info("👤 Create your bank account")

        with col2:
            st.info("💰 Manage your money")

        with col3:
            st.info("📋 Track your transactions")

    elif main_option == "Create Account":

        create_account()

    elif main_option == "Login":

        login()

else:

    sub_menu(
        st.session_state.account_number
    )

